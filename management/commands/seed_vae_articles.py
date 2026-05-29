import os
import random
import tempfile

from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management.base import BaseCommand, CommandError
from django.utils import timezone

from core import files as core_files
from core.models import Account
from journal.models import Journal
from submission import models as submission_models

from plugins.vae_workflow import models, plugin_settings


SAMPLE_TITLES = [
    'A Phenomenological Study of Late Romantic Aesthetics',
    'Diasporic Memory in Twentieth-Century Travel Writing',
    'Computational Approaches to Stylistic Variation in Middle English',
    'Reimagining the Archive: Open Access and Cultural Heritage',
    'Performing Identity in Postcolonial Theatre',
    'Music, Mediation, and the Public Sphere',
    'Translation as Ethical Practice in Comparative Literature',
    'Material Texts and Reading Communities',
    'Cosmopolitanism and the Limits of the Nation-State',
    'Voice, Authorship, and Authority in Digital Media',
]

SAMPLE_ABSTRACT = (
    'This article examines the interplay of cultural form and historical '
    'contingency through a series of close readings. Drawing on archival '
    'sources and recent theoretical work, it argues that critical attention '
    'to the margins of canonical practice yields fresh insight into the '
    'period under discussion. The argument unfolds in three parts: first, a '
    'reconstruction of the relevant historiography; second, a series of '
    'case studies that test the proposed framework; and third, a synthesis '
    'that points toward further avenues for research.'
)

SAMPLE_KEYWORDS = [
    'cultural history', 'critical theory', 'archival research',
    'open access', 'digital humanities', 'comparative literature',
    'reception studies', 'aesthetics', 'historiography',
]

AUTHOR_NAMES = [
    ('Ada', 'Lovelace'), ('Bertrand', 'Russell'), ('Cordelia', 'Fine'),
    ('Donna', 'Haraway'), ('Edward', 'Said'), ('Frances', 'Yates'),
    ('Gayatri', 'Spivak'), ('Hannah', 'Arendt'), ('Isaiah', 'Berlin'),
    ('Judith', 'Butler'),
]


def build_docx(title, abstract):
    try:
        from docx import Document
    except ImportError:
        raise CommandError(
            'python-docx is required. Install with `pip install python-docx`.'
        )
    doc = Document()
    doc.add_heading(title, level=1)
    doc.add_paragraph(abstract)
    doc.add_heading('Introduction', level=2)
    doc.add_paragraph(
        'This is a synthetic test manuscript generated for the VAE workflow. '
        'It contains placeholder text intended only for development and '
        'testing purposes.'
    )
    doc.add_heading('Method', level=2)
    doc.add_paragraph(
        'The methodology is illustrative. Real submissions will replace this '
        'content with substantive scholarly material.'
    )
    doc.add_heading('Discussion', level=2)
    doc.add_paragraph(
        'Discussion proceeds in line with established conventions in the '
        'field. Citations and references would be elaborated in a complete '
        'manuscript.'
    )
    fd, path = tempfile.mkstemp(suffix='.docx')
    os.close(fd)
    doc.save(path)
    return path


def get_or_create_author(journal, index):
    first, last = AUTHOR_NAMES[index % len(AUTHOR_NAMES)]
    email = '{}.{}.{}@example.com'.format(
        first.lower(), last.lower(), index,
    )
    account = Account.objects.filter(email=email).first()
    if not account:
        account = Account.objects.create_user(
            email=email,
            password='test1234',
            first_name=first,
            last_name=last,
        )
        account.is_active = True
        account.institution = 'University of Example'
        account.save()
    return account


class Command(BaseCommand):
    help = 'Create N test articles in the VAE Claiming stage of a journal.'

    def add_arguments(self, parser):
        parser.add_argument(
            'count',
            type=int,
            help='Number of articles to create.',
        )
        parser.add_argument(
            '--journal',
            default='olh',
            help='Journal code (default: olh).',
        )

    def handle(self, *args, **options):
        count = options['count']
        journal_code = options['journal']

        try:
            journal = Journal.objects.get(code=journal_code)
        except Journal.DoesNotExist:
            raise CommandError(
                'Journal with code "{}" not found.'.format(journal_code)
            )

        section = submission_models.Section.objects.filter(
            journal=journal,
        ).first()
        if not section:
            section = submission_models.Section.objects.create(
                journal=journal,
                name='Articles',
            )

        licence = submission_models.Licence.objects.filter(
            journal=journal,
        ).first()

        language_options = submission_models.LANGUAGE_CHOICES
        default_language = 'eng' if any(
            code == 'eng' for code, _ in language_options
        ) else language_options[0][0]

        created = []
        for i in range(count):
            title = '{} ({})'.format(
                SAMPLE_TITLES[i % len(SAMPLE_TITLES)],
                timezone.now().strftime('%Y-%m-%d %H:%M:%S'),
            )
            author = get_or_create_author(journal, i)

            article = submission_models.Article.objects.create(
                journal=journal,
                title=title,
                abstract=SAMPLE_ABSTRACT,
                language=default_language,
                section=section,
                license=licence,
                stage=plugin_settings.STAGE,
                date_submitted=timezone.now(),
                owner=author,
                correspondence_author=author,
            )

            keyword_sample = random.sample(SAMPLE_KEYWORDS, k=3)
            for word in keyword_sample:
                keyword, _ = submission_models.Keyword.objects.get_or_create(
                    word=word,
                )
                article.keywords.add(keyword)

            article.authors.add(author)
            try:
                from submission.models import FrozenAuthor
                FrozenAuthor.objects.create(
                    article=article,
                    author=author,
                    first_name=author.first_name,
                    last_name=author.last_name,
                    email=author.email,
                    institution=author.institution or 'University of Example',
                    order=0,
                )
            except Exception:
                pass

            docx_path = build_docx(title, SAMPLE_ABSTRACT)
            try:
                with open(docx_path, 'rb') as fh:
                    uploaded = SimpleUploadedFile(
                        name='{}.docx'.format(
                            title.split(' ')[0].lower()
                        ),
                        content=fh.read(),
                        content_type=(
                            'application/vnd.openxmlformats-officedocument'
                            '.wordprocessingml.document'
                        ),
                    )
                file_obj = core_files.save_file_to_article(
                    uploaded,
                    article,
                    owner=author,
                    label='Manuscript',
                    description='Synthetic test manuscript.',
                )
                article.manuscript_files.add(file_obj)
            finally:
                if os.path.exists(docx_path):
                    os.remove(docx_path)

            models.ArticlePoolAvailability.objects.get_or_create(
                article=article,
                defaults={'available': False},
            )

            created.append(article)
            self.stdout.write(
                self.style.SUCCESS(
                    'Created article #{} — {}'.format(article.pk, title)
                )
            )

        self.stdout.write(
            self.style.SUCCESS(
                '\nCreated {} article(s) in journal "{}" at stage "{}".'.format(
                    len(created), journal.code, plugin_settings.STAGE,
                )
            )
        )
